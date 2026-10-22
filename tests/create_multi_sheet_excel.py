import csv
from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parent.parent
CLASSIFIED_DIR = ROOT / "outputs" / "classified"
EXCEL_PATH = ROOT / "outputs" / "模型发售与再版分类整理表.xlsx"

# 11 个分类定义（包含用户已手动更新的 STAR-星动.csv）
SHEET_CONFIG = [
    ("STAR-星动", CLASSIFIED_DIR / "STAR-星动.csv"),
    ("高高", CLASSIFIED_DIR / "高高.csv"),
    ("长龙", CLASSIFIED_DIR / "长龙.csv"),
    ("SUPA", CLASSIFIED_DIR / "SUPA.csv"),
    ("四海", CLASSIFIED_DIR / "四海.csv"),
    ("GL", CLASSIFIED_DIR / "GL.csv"),
    ("星聚模型", CLASSIFIED_DIR / "星聚模型.csv"),
    ("千叶模型", CLASSIFIED_DIR / "千叶模型.csv"),
    ("神兽模型", CLASSIFIED_DIR / "神兽模型.csv"),
    ("乐豹模型", CLASSIFIED_DIR / "乐豹模型.csv"),
    ("其他与小众厂商", CLASSIFIED_DIR / "其他与小众厂商.csv"),
]

wb = Workbook()
wb.remove(wb.active) # 移除默认空白 sheet

# 样式定义
font_header = Font(name="微软雅黑", size=11, bold=True, color="FFFFFF")
fill_header = PatternFill(start_color="2F5597", end_color="2F5597", fill_type="solid")
fill_master_header = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")

font_data = Font(name="微软雅黑", size=10)
fill_alt = PatternFill(start_color="F7F9FC", end_color="F7F9FC", fill_type="solid")

thin_border = Border(
    left=Side(style="thin", color="D9D9D9"),
    right=Side(style="thin", color="D9D9D9"),
    top=Side(style="thin", color="D9D9D9"),
    bottom=Side(style="thin", color="D9D9D9")
)

align_center = Alignment(horizontal="center", vertical="center")
align_left = Alignment(horizontal="left", vertical="center")

all_rows_for_summary = []
sheet_counts = {}

# 遍历各分类 CSV 生成独立 Sheet
for sheet_name, csv_file in SHEET_CONFIG:
    ws = wb.create_sheet(title=sheet_name)
    ws.views.sheetView[0].showGridLines = True
    ws.freeze_panes = "A2"
    
    with open(csv_file, "r", encoding="utf-8-sig") as f:
        reader = csv.reader(f)
        header = next(reader)
        ws.append(header)
        
        # 格式化表头
        for col_num in range(1, len(header) + 1):
            cell = ws.cell(row=1, column=col_num)
            cell.font = font_header
            cell.fill = fill_header
            cell.alignment = align_center
            cell.border = thin_border
        
        # 写入有效数据行
        row_idx = 2
        count = 0
        for r in reader:
            # 过滤空行或仅含逗号的残余行
            if not any(cell.strip() for cell in r):
                continue
            ws.append(r)
            count += 1
            all_rows_for_summary.append((r[0], r[1], r[2], r[3] if len(r) > 3 else "", sheet_name))
            
            is_alt = (row_idx % 2 == 0)
            for col_num in range(1, len(r) + 1):
                cell = ws.cell(row=row_idx, column=col_num)
                cell.font = font_data
                cell.border = thin_border
                if is_alt:
                    cell.fill = fill_alt
                if col_num in (1, 2):
                    cell.alignment = align_center
                else:
                    cell.alignment = align_left
            row_idx += 1
        sheet_counts[sheet_name] = count

    # 自适应列宽
    for col in ws.columns:
        max_len = 0
        for cell in col:
            val_str = str(cell.value or "")
            length = sum(2 if ord(c) > 127 else 1 for c in val_str)
            if length > max_len:
                max_len = length
        col_letter = get_column_letter(col[0].column)
        ws.column_dimensions[col_letter].width = max(max_len + 4, 12)

# 新增「全部记录汇总」放置在第 1 个 Sheet，方便全量筛选
ws_all = wb.create_sheet(title="全部记录汇总", index=0)
ws_all.views.sheetView[0].showGridLines = True
ws_all.freeze_panes = "A2"

summary_header = ["日期", "厂家", "模型", "原始记录", "所属分类Sheet"]
ws_all.append(summary_header)
for col_num in range(1, len(summary_header) + 1):
    cell = ws_all.cell(row=1, column=col_num)
    cell.font = font_header
    cell.fill = fill_master_header
    cell.alignment = align_center
    cell.border = thin_border

for row_idx, r in enumerate(all_rows_for_summary, start=2):
    ws_all.append(r)
    is_alt = (row_idx % 2 == 0)
    for col_num in range(1, len(r) + 1):
        cell = ws_all.cell(row=row_idx, column=col_num)
        cell.font = font_data
        cell.border = thin_border
        if is_alt:
            cell.fill = fill_alt
        if col_num in (1, 2, 5):
            cell.alignment = align_center
        else:
            cell.alignment = align_left

for col in ws_all.columns:
    max_len = 0
    for cell in col:
        val_str = str(cell.value or "")
        length = sum(2 if ord(c) > 127 else 1 for c in val_str)
        if length > max_len:
            max_len = length
    col_letter = get_column_letter(col[0].column)
    ws_all.column_dimensions[col_letter].width = max(max_len + 4, 12)

wb.save(EXCEL_PATH)
print(f"Workbook generated successfully at: {EXCEL_PATH}")
print(f"Total sheets: {len(wb.sheetnames)}")
for sname, cnt in sheet_counts.items():
    print(f"  - {sname:15s}: {cnt:3d} 条")
print(f"Total in '全部记录汇总': {len(all_rows_for_summary)} 条")
