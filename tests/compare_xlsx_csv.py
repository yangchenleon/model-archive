import openpyxl
import csv

wb = openpyxl.load_workbook("outputs/模型发售与再版分类整理表.xlsx", data_only=True)
star_rows = list(wb['STAR-星动'].iter_rows(values_only=True))[1:]
supa_rows = list(wb['SUPA'].iter_rows(values_only=True))[1:]

with open("outputs/星动系_全量整理.csv", "r", encoding="utf-8") as f:
    reader = csv.DictReader(f)
    field_vendor = reader.fieldnames[0]
    csv_rows = list(reader)

print(f"Total rows in STAR-星动 sheet: {len(star_rows)}")
print(f"Total rows in SUPA sheet: {len(supa_rows)}")
print(f"Total rows in 星动系_全量整理.csv: {len(csv_rows)}")

# Let's inspect how items match or don't match
print("\n--- Check SUPA sheet items against CSV ---")
for r in supa_rows:
    date, mfg, name, raw = r
    # Find potential matches in CSV
    matches = [c for c in csv_rows if name in c['模型名称'] or (c['版本/配色'] and c['版本/配色'] in name) or name in c['详情'] or (c['模型名称'] in name)]
    print(f"SUPA: {date} | {name} => matched {len(matches)} rows")
    for m in matches[:2]:
        print(f"   CSV: [{m[field_vendor]}][{m['产品线']}] {m['模型名称']} | {m['版本/配色']} | {m['详情'][:30]}")

print("\n--- Check STAR-星动 sheet items against CSV ---")
matched_count = 0
unmatched = []
for r in star_rows:
    date, mfg, name, raw = r
    matches = [c for c in csv_rows if name in c['模型名称'] or (c['模型名称'] in name) or name in c['详情']]
    if matches:
        matched_count += 1
    else:
        unmatched.append(r)

print(f"STAR-星动 sheet: {matched_count} matched, {len(unmatched)} not directly matched")
print("\nUnmatched samples (first 20):")
for u in unmatched[:20]:
    print(u)
